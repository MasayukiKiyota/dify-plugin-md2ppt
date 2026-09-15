"""SDK 非依存のスモークテスト。

    .venv/Scripts/python.exe tests/test_convert.py

dify_plugin をインストールしなくても変換パイプライン全体を検証できる。
"""

from __future__ import annotations

import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import md2ppt_utils as u  # noqa: E402

FIXTURES = Path(__file__).resolve().parent / "fixtures"
SAMPLE_MD = FIXTURES / "sample.md"
TEMPLATE = FIXTURES / "template.pptx"
CONFIG = FIXTURES / "config.yaml"

_POTX_CT = ("application/vnd.openxmlformats-officedocument"
            ".presentationml.template.main+xml")
_PPTX_CT = ("application/vnd.openxmlformats-officedocument"
            ".presentationml.presentation.main+xml")

failures: list[str] = []


def _has_marl(pptx_bytes: bytes) -> bool:
    """段落に marL（明示インデント）が書かれているか。

    表のセル余白も <a:tcPr marL="..."> を持つので、段落の <a:pPr> だけを見る。
    """
    import io
    import re

    with zipfile.ZipFile(io.BytesIO(pptx_bytes)) as z:
        return any(
            re.search(r'<a:pPr[^>]* marL="', z.read(n).decode("utf-8"))
            for n in z.namelist()
            if n.startswith("ppt/slides/slide")
        )


def _rpr_tags(pptx_bytes: bytes, *prefixes: str) -> list[str]:
    """lang / noProof を持てる要素の開始タグを集める。

    prefixes を渡すとそのパートだけに絞る（"ppt/slides/" など）。
    """
    import io
    import re

    found: list[str] = []
    with zipfile.ZipFile(io.BytesIO(pptx_bytes)) as z:
        for name in z.namelist():
            if not name.endswith(".xml"):
                continue
            if prefixes and not name.startswith(prefixes):
                continue
            xml = z.read(name).decode("utf-8", "ignore")
            found += re.findall(r"<a:(?:rPr|endParaRPr|defRPr)\b[^>]*>", xml)
    return found


def _attr_values(tags: list[str], attr: str) -> set[str]:
    import re

    pattern = re.compile(r'\b%s="([^"]*)"' % attr)
    return {m.group(1) for m in map(pattern.search, tags) if m}


def langs(pptx_bytes: bytes, *prefixes: str) -> set[str]:
    """書き込まれている lang の値。"""
    return _attr_values(_rpr_tags(pptx_bytes, *prefixes), "lang")


def noproofs(pptx_bytes: bytes, *prefixes: str) -> set[str]:
    """書き込まれている noProof の値。"""
    return _attr_values(_rpr_tags(pptx_bytes, *prefixes), "noProof")


def unlabeled_rprs(pptx_bytes: bytes) -> int:
    """lang を持たない rPr / endParaRPr / defRPr の数。0 が理想。"""
    import re

    return sum(1 for t in _rpr_tags(pptx_bytes) if not re.search(r'\blang="', t))


def strip_langs(pptx_bytes: bytes) -> bytes:
    """lang 属性を全部落とした .pptx を作る（フォールバック検証用）。

    altLang は残す（直前が "t" なので \slang= には一致しない）。
    """
    import io
    import re

    src = io.BytesIO(pptx_bytes)
    out = io.BytesIO()
    with zipfile.ZipFile(src) as zf:
        entries = [(i, zf.read(i.filename)) for i in zf.infolist()]
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as w:
        for info, data in entries:
            if info.filename.endswith(".xml"):
                data = re.sub(r'\slang="[^"]*"', "",
                              data.decode("utf-8", "ignore")).encode("utf-8")
            w.writestr(info, data)
    return out.getvalue()


def run_sizes(pptx_bytes: bytes) -> set[float]:
    """スライドの run に書かれた文字サイズ（pt）。

    a:defRPr / a:endParaRPr は継承の定義なので数えない。実際に描かれる
    テキストが持つ <a:rPr sz="..."> だけを見る。
    """
    import io
    import re

    found: set[float] = set()
    with zipfile.ZipFile(io.BytesIO(pptx_bytes)) as z:
        for name in z.namelist():
            if not name.startswith("ppt/slides/slide"):
                continue
            xml = z.read(name).decode("utf-8", "ignore")
            for tag in re.findall(r"<a:rPr\b[^>]*>", xml):
                m = re.search(r'\bsz="(\d+)"', tag)
                if m:
                    found.add(int(m.group(1)) / 100.0)
    return found


def unsized_runs(pptx_bytes: bytes) -> int:
    """sz を持たない run の数（＝テンプレートの文字サイズを継承する run）。"""
    import io
    import re

    total = 0
    with zipfile.ZipFile(io.BytesIO(pptx_bytes)) as z:
        for name in z.namelist():
            if not name.startswith("ppt/slides/slide"):
                continue
            xml = z.read(name).decode("utf-8", "ignore")
            total += sum(1 for t in re.findall(r"<a:rPr\b[^>]*>", xml)
                         if not re.search(r'\bsz="', t))
    return total


def strip_sizes(pptx_bytes: bytes) -> bytes:
    """sz 属性を全部落とした .pptx を作る（フォールバック検証用）。"""
    import io
    import re

    src = io.BytesIO(pptx_bytes)
    out = io.BytesIO()
    with zipfile.ZipFile(src) as zf:
        entries = [(i, zf.read(i.filename)) for i in zf.infolist()]
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as w:
        for info, data in entries:
            if info.filename.endswith(".xml"):
                data = re.sub(r'\ssz="[^"]*"', "",
                              data.decode("utf-8", "ignore")).encode("utf-8")
            w.writestr(info, data)
    return out.getvalue()


def check(label: str, cond: bool, detail: str = "") -> None:
    if cond:
        print(f"  ok   {label}")
    else:
        print(f"  FAIL {label} {detail}")
        failures.append(label)


def as_potx(pptx_bytes: bytes) -> bytes:
    """.pptx を .potx 相当（content type だけ template）に作り替える。"""
    import io
    src = io.BytesIO(pptx_bytes)
    out = io.BytesIO()
    with zipfile.ZipFile(src) as zf:
        entries = [(i, zf.read(i.filename)) for i in zf.infolist()]
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as w:
        for info, data in entries:
            if info.filename == "[Content_Types].xml":
                data = data.decode("utf-8").replace(_PPTX_CT, _POTX_CT).encode("utf-8")
            w.writestr(info, data)
    return out.getvalue()


def english_template(anonymize: bool = False) -> bytes:
    """python-pptx 既定のテンプレート（標準的な英語レイアウト 11 種）。

    anonymize=True でレイアウト名を Layout-NN に潰し、名前のヒントを完全に消す。
    """
    import io

    from pptx import Presentation

    prs = Presentation()
    if anonymize:
        for i, layout in enumerate(prs.slide_layouts):
            layout.name = f"Layout-{i:02d}"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def main() -> int:
    md = SAMPLE_MD.read_text(encoding="utf-8")
    tpl = TEMPLATE.read_bytes()
    cfg = CONFIG.read_text(encoding="utf-8")

    print("[1] convert with sample config")
    r = u.convert(md, tpl, "template.pptx", cfg)
    check("returns a zip", r["pptx"][:2] == b"PK")
    check("slide_count >= spec_count",
          r["slide_count"] >= r["spec_count"],
          f"({r['slide_count']} vs {r['spec_count']})")
    check("no warnings", not r["warnings"], str(r["warnings"]))
    check("front matter title picked up", r["meta"].get("title") is not None,
          str(r["meta"]))
    check("outline is non-empty", bool(r["outline"].strip()))
    print(f"       slides={r['slide_count']} specs={r['spec_count']} "
          f"bytes={len(r['pptx'])} meta={r['meta']}")
    print("       " + r["outline"].replace("\n", "\n       "))

    print("[2] convert with DEFAULT_CONFIG (config_yaml 未指定)")
    r2 = u.convert(md, tpl, "template.pptx", None)
    check("same slide count as sample config",
          r2["slide_count"] == r["slide_count"],
          f"({r2['slide_count']} vs {r['slide_count']})")
    check("no warnings", not r2["warnings"], str(r2["warnings"]))

    print("[3] .potx (content type = template) is normalized")
    potx = as_potx(tpl)
    r3 = u.convert(md, potx, "corporate.potx", cfg)
    check("potx converts", r3["pptx"][:2] == b"PK")
    check("potx slide count matches pptx",
          r3["slide_count"] == r["slide_count"])

    print("[4] analyze_template")
    info = u.analyze_template(tpl, "template.pptx")
    check("11 layouts", info["layout_count"] == 11, str(info["layout_count"]))
    check("16:9", info["aspect_ratio"] == "16:9", info["aspect_ratio"])
    check("13.33 x 7.5 in",
          (info["slide_width_in"], info["slide_height_in"]) == (13.33, 7.5),
          f"{info['slide_width_in']} x {info['slide_height_in']}")
    check("表紙 layout present", "表紙" in info["layout_names"],
          str(info["layout_names"]))
    report = u.format_template_report(info)
    check("report mentions every layout",
          all(n in report for n in info["layout_names"]))

    print("[5] detect_config round-trips through convert")
    sug = u.detect_config(info)
    check("guessed 表紙/章扉/本文",
          (sug["layouts"]["title"], sug["layouts"]["section"],
           sug["layouts"]["content"]) == ("表紙", "章扉", "本文"),
          str(sug["layouts"]))
    yml = u.config_to_yaml(sug, "template.pptx")
    r5 = u.convert(md, tpl, "template.pptx", yml)
    check("suggested config converts cleanly", not r5["warnings"], str(r5["warnings"]))

    print("[6] analyze_template on .potx")
    info_p = u.analyze_template(potx, "corporate.potx")
    check("potx analyzed", info_p["layout_count"] == 11)

    print("[7] error paths")
    for label, fn in [
        ("bad extension", lambda: u.convert(md, tpl, "deck.ppt", None)),
        ("not a zip", lambda: u.convert(md, b"hello world", "t.pptx", None)),
        ("empty markdown", lambda: u.convert("   ", tpl, "t.pptx", None)),
        ("bad yaml", lambda: u.convert(md, tpl, "t.pptx", "layouts: [oops")),
        # コロンを含むのでマッピングを書こうとした形跡がある = 握りつぶさない
        ("list yaml", lambda: u.convert(md, tpl, "t.pptx", "- layouts: 本文")),
        ("multi-line scalar",
         lambda: u.convert(md, tpl, "t.pptx", "設定を書き忘れました\nここには何もありません\n")),
    ]:
        try:
            fn()
            check(f"{label} raises Md2pptError", False, "(no exception)")
        except u.Md2pptError as e:
            check(f"{label} raises Md2pptError", True)
            print(f"       -> {e}")
        except Exception as e:
            check(f"{label} raises Md2pptError", False,
                  f"(got {type(e).__name__}: {e})")

    print("[7b] 見出しの無い Markdown は本文だけの 1 枚になる")
    r7 = u.convert("plain text only", tpl, "t.pptx", None)
    check("one untitled slide", r7["spec_count"] == 1, str(r7["spec_count"]))

    print("[8] unknown layout names produce warnings, not a crash")
    r8 = u.convert(md, tpl, "t.pptx", 'layouts:\n  content: "存在しないレイアウト"\n')
    check("still converts", r8["pptx"][:2] == b"PK")
    check("warns about the missing layout", any("存在しないレイアウト" in w
                                                for w in r8["warnings"]),
          str(r8["warnings"]))

    print("[9] input helpers")
    check("ensure_pptx appends suffix", u.ensure_pptx("deck") == "deck.pptx")
    check("ensure_pptx sanitizes", u.ensure_pptx("a/b:c.pptx") == "a_b_c.pptx",
          u.ensure_pptx("a/b:c.pptx"))
    check("ensure_pptx falls back", u.ensure_pptx("  ") == "presentation.pptx")
    check("decode cp932", u.decode_text("日本語".encode("cp932")) == "日本語")
    check("decode utf-8 BOM", u.decode_text("# あ".encode("utf-8-sig")) == "# あ")
    # text_param が前後の空白を落とすので、末尾の改行は残らない
    text, used = u.read_markdown("  # a\r\nb\r\n  ", None)
    check("CRLF normalized", text == "# a\nb", repr(text))
    check("used_file False for text", used is False)

    class FakeFile:
        blob = "# ファイル入力\n".encode("utf-8")
        filename = "in.md"

    text, used = u.read_markdown(None, FakeFile())
    check("File param decoded", text.startswith("# ファイル入力"))
    check("used_file True for file", used is True)
    check("file_name_of", u.file_name_of(FakeFile()) == "in.md")
    check("file_bytes(None) is None", u.file_bytes(None) is None)

    print("[10] config_yaml 無しで名前の違うテンプレートに追従する")
    en = english_template()
    r10 = u.convert(md, en, "default.pptx", None)
    check("english template needs no config", not r10["warnings"], str(r10["warnings"]))
    check("english layouts detected by name",
          r10["layouts_used"] == {
              "title": "Title Slide", "section": "Section Header",
              "content": "Title and Content", "table": "Title Only",
              "blank": "Blank"},
          str(r10["layouts_used"]))
    check("no extra slides from bad layouts",
          r10["slide_count"] == r10["spec_count"] == 13,
          f"{r10['slide_count']}/{r10['spec_count']}")

    print("[11] 名前が手がかりにならなくても構造から判定する")
    anon = english_template(anonymize=True)
    info_a = u.analyze_template(anon, "anon.pptx")
    check("names carry no hint",
          all(n.startswith("Layout-") for n in info_a["layout_names"]),
          str(info_a["layout_names"]))
    r11 = u.convert(md, anon, "anon.pptx", None)
    check("anonymous template needs no config", not r11["warnings"], str(r11["warnings"]))
    check("cover = the layout with a SUBTITLE",
          r11["layouts_used"]["title"] == "Layout-00", str(r11["layouts_used"]))
    check("content = title + exactly one body",
          r11["layouts_used"]["content"] == "Layout-01", str(r11["layouts_used"]))
    check("table = title with no body",
          r11["layouts_used"]["table"] == "Layout-05", str(r11["layouts_used"]))
    check("blank = no placeholders at all",
          r11["layouts_used"]["blank"] == "Layout-06", str(r11["layouts_used"]))
    check("section falls back to the title-only layout",
          r11["layouts_used"]["section"] == "Layout-05", str(r11["layouts_used"]))

    print("[12] config_yaml は自動判定より優先される")
    r12 = u.convert(md, anon, "anon.pptx", 'layouts:\n  content: "Layout-07"\n')
    check("explicit content layout wins",
          r12["layouts_used"]["content"] == "Layout-07", str(r12["layouts_used"]))
    check("unspecified keys keep the detected value",
          r12["layouts_used"]["title"] == "Layout-00", str(r12["layouts_used"]))

    print("[13] 既定テンプレートの判定は変わらない（回帰確認）")
    det = u.detect_config(u.analyze_template(tpl, "template.pptx"))
    check("japanese template still detected by name",
          det["layouts"] == {"title": "表紙", "section": "章扉", "content": "本文",
                             "table": "タイトルのみ", "blank": "白紙"},
          str(det["layouts"]))
    check("body placeholder idx detected", det["placeholders"]["body"] == 1,
          str(det["placeholders"]))

    print("[14] 未入力の任意パラメータが番兵値で届いても未指定として扱う")
    # Dify は未入力の任意パラメータを str(None) 経由で "None" として送ってくることがある。
    for sentinel in ["None", "none", "NONE", "null", "~", "undefined", "nil",
                     "-", "", "   ", None]:
        try:
            rs = u.convert(md, tpl, "t.pptx", sentinel)
            ok = not rs["warnings"] and rs["slide_count"] == 13
        except Exception as e:
            ok = False
            print(f"       {sentinel!r} -> {type(e).__name__}: {e}")
        check(f"config_yaml={sentinel!r} は未指定扱い", ok)
    check("ensure_pptx('None') が None.pptx にならない",
          u.ensure_pptx("None") == "presentation.pptx", u.ensure_pptx("None"))
    check("ensure_pptx('-') も既定名", u.ensure_pptx("-") == "presentation.pptx")
    for sentinel in ["None", "null", "-"]:
        try:
            u.read_markdown(sentinel, None)
            check(f"markdown_text={sentinel!r} は空扱い", False, "(no exception)")
        except u.Md2pptError:
            check(f"markdown_text={sentinel!r} は空扱い", True)

    print("[14b] 安全弁: コロンも改行も無い単一トークンは未入力扱い")
    # 番兵値の列挙に無い未知の文字列が届いても変換を止めない
    for token in ["None", "undefined", "xyzzy", "設定なし", "false", "0", "[]"]:
        try:
            rt = u.convert(md, tpl, "t.pptx", token)
            ok = not rt["warnings"] and rt["slide_count"] == 13
        except Exception as e:
            ok = False
            print(f"       {token!r} -> {type(e).__name__}: {e}")
        check(f"config_yaml={token!r} は未入力扱い", ok)
    check("コロンがあれば握りつぶさない",
          u.parse_config_yaml("layouts: {content: 本文}") == {"layouts": {"content": "本文"}})

    print("[15] 全角文字は具体的に指摘される")
    # 全角コロンは YAML の構文エラーにならず、設定全体がただの文字列になる
    fullwidth_colon = """layouts：
  content："本文"
"""
    # 全角スペース字下げは {'layouts': None, '　　content': ...} を作ってしまう
    fullwidth_space = """layouts:
　　content: "本文"
"""
    for label, bad, want in [("全角コロン", fullwidth_colon, "全角コロン"),
                             ("全角スペース", fullwidth_space, "全角スペース")]:
        try:
            u.convert(md, tpl, "t.pptx", bad)
            check(f"{label}を指摘", False, "(no exception)")
        except u.Md2pptError as e:
            check(f"{label}を指摘", want in str(e), str(e))
            print(f"       -> {e}")

    print("[16] bool_param は文字列の真偽値を正しく解釈する")
    for value, expected in [("false", False), ("False", False), ("true", True),
                            ("True", True), ("0", False), ("1", True),
                            ("None", False), ("", False), (None, False),
                            (False, False), (True, True), ("off", False)]:
        check(f"bool_param({value!r}) is {expected}",
              u.bool_param(value) is expected, str(u.bool_param(value)))

    print("[17] 設定の形がおかしいときは分かりやすく落ちる")
    unknown_key = """layout:
  content: "本文"
"""
    scalar_section = """layouts: 本文
"""
    null_section = """layouts:
placeholders:
  body: 1
"""
    for label, bad in [("不明なトップレベルキー", unknown_key),
                       ("セクションがマッピングでない", scalar_section),
                       ("セクションが null", null_section)]:
        try:
            u.convert(md, tpl, "t.pptx", bad)
            check(f"{label} はエラー", False, "(no exception)")
        except u.Md2pptError as e:
            check(f"{label} はエラー", True)
            print(f"       -> {e}")

    with_template = """template: foo.pptx
layouts:
  content: "本文"
"""
    check("template キーは黙って無視される",
          u.parse_config_yaml(with_template)
          == {"template": "foo.pptx", "layouts": {"content": "本文"}})
    check("template キーがあっても変換できる",
          u.convert(md, tpl, "t.pptx", with_template)["warnings"] == [])

    print("[18] fonts を null にするとフォントを書き込まない")

    def typefaces(pptx_bytes: bytes) -> set[str]:
        """生成された .pptx の run に書かれた typeface を全部集める。"""
        import io

        from pptx import Presentation
        from pptx.util import Emu  # noqa: F401  (pptx の遅延 import を促す)

        ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
        found: set[str] = set()
        prs = Presentation(io.BytesIO(pptx_bytes))
        for slide in prs.slides:
            for tag in ("latin", "ea", "cs"):
                for el in slide.shapes._spTree.iter(f"{ns}{tag}"):
                    found.add(el.get("typeface"))
        return found

    no_fonts = """fonts:
  latin: null
  eastasian: null
  code: null
  code_eastasian: null
"""
    r18 = u.convert(md, tpl, "t.pptx", no_fonts)
    check("null fonts converts", not r18["warnings"], str(r18["warnings"]))
    check("同じ枚数になる", r18["slide_count"] == r["slide_count"],
          f"({r18['slide_count']} vs {r['slide_count']})")
    check("typeface が 1 つも書かれない", typefaces(r18["pptx"]) == set(),
          str(typefaces(r18["pptx"])))
    check("設定なし（既定）でも typeface は書かれない", typefaces(r2["pptx"]) == set(),
          str(sorted(typefaces(r2["pptx"]))))
    check("明示指定すれば typeface が書かれる", "Calibri" in typefaces(r["pptx"]),
          str(sorted(typefaces(r["pptx"]))))

    print("[18b] 一部だけ null にもできる")
    ea_only = """fonts:
  latin: null
  eastasian: Meiryo
  code: null
  code_eastasian: null
"""
    r18b = u.convert(md, tpl, "t.pptx", ea_only)
    check("片側だけ null で変換できる", not r18b["warnings"], str(r18b["warnings"]))
    check("指定した書体だけが書かれる", typefaces(r18b["pptx"]) == {"Meiryo"},
          str(sorted(typefaces(r18b["pptx"]))))
    # 空文字も「指定しない」として扱う（Dify から空欄が "" で届くことがある）
    blank_fonts = """fonts:
  latin: ""
  eastasian: ""
  code: ""
  code_eastasian: ""
"""
    r18c = u.convert(md, tpl, "t.pptx", blank_fonts)
    check("空文字も未指定扱い", typefaces(r18c["pptx"]) == set(),
          str(sorted(typefaces(r18c["pptx"]))))

    print("[19] colors を null にすると色を塗らない")

    def srgb(pptx_bytes: bytes) -> set[str]:
        """スライドに書かれた固定色（srgbClr）を全部集める。"""
        import io
        import re

        found: set[str] = set()
        with zipfile.ZipFile(io.BytesIO(pptx_bytes)) as z:
            for name in z.namelist():
                if name.startswith("ppt/slides/slide"):
                    found |= set(re.findall(r'srgbClr val="([0-9A-Fa-f]{6})"',
                                            z.read(name).decode("utf-8")))
        return found

    # fixtures/config.yaml は 14 色すべてを明示している。
    explicit = u.parse_config_yaml(CONFIG.read_text(encoding="utf-8"))["colors"]
    color_keys = sorted(u.core.DEFAULT_CONFIG["colors"])
    check("既定値は 14 色すべて null",
          set(u.core.DEFAULT_CONFIG["colors"].values()) == {None},
          str(u.core.DEFAULT_CONFIG["colors"]))
    check("fixtures/config.yaml は 14 色すべてを明示",
          sorted(explicit) == color_keys, str(sorted(explicit)))

    def colors_yaml(overrides: dict) -> str:
        merged = {**explicit, **overrides}
        body = "".join(
            f"  {k}: " + ("null" if v is None else f'"{v}"') + "\n"
            for k, v in merged.items()
        )
        # 表のセルを塗る経路（fill_cell）も通したいので明示的に有効にする。
        return "colors:\n" + body + "table:\n  explicit_format: true\n"

    check("設定なし（既定）では固定色が書かれない", srgb(r2["pptx"]) == set(),
          str(sorted(srgb(r2["pptx"]))))
    check("明示指定すれば固定色が書かれる",
          set(explicit.values()) <= srgb(r["pptx"]),
          str(sorted(set(explicit.values()) - srgb(r["pptx"]))))

    no_colors = colors_yaml({k: None for k in color_keys})
    r19 = u.convert(md, tpl, "t.pptx", no_colors)
    check("null colors converts", not r19["warnings"], str(r19["warnings"]))
    check("同じ枚数になる", r19["slide_count"] == r["slide_count"],
          f"({r19['slide_count']} vs {r['slide_count']})")
    check("固定色が 1 つも残らない",
          not (srgb(r19["pptx"]) & set(explicit.values())),
          str(sorted(srgb(r19["pptx"]))))

    print("[19b] 一部だけ null にもできる")
    # rgb() を直接呼んでいた図形（コード枠・引用バー・罫線・表）も落ちないこと。
    for key in color_keys:
        try:
            one = u.convert(md, tpl, "t.pptx", colors_yaml({key: None}))
            # 同じ色を使う別のキーが残っていれば、その色は消えなくてよい。
            shared = any(v == explicit[key] for k, v in explicit.items() if k != key)
            check(f"{key}: null で色が消える",
                  shared or explicit[key] not in srgb(one["pptx"]),
                  str(sorted(srgb(one["pptx"]))))
        except Exception as e:                                  # noqa: BLE001
            check(f"{key}: null で色が消える", False, f"{type(e).__name__}: {e}")

    blank_colors = "colors:\n" + "".join(f'  {k}: ""\n' for k in color_keys)
    r19c = u.convert(md, tpl, "t.pptx", blank_colors)
    check("空文字も未指定扱い",
          not (srgb(r19c["pptx"]) & set(explicit.values())),
          str(sorted(srgb(r19c["pptx"]))))

    r19d = u.convert(md, tpl, "t.pptx", no_fonts + no_colors)
    check("fonts と colors の null を併用できる",
          typefaces(r19d["pptx"]) == set()
          and not (srgb(r19d["pptx"]) & set(explicit.values())))

    print("[19c] 既定では表スタイル ID と箇条書きインデントも指定しない")
    check("table.style_id の既定は null",
          u.core.DEFAULT_CONFIG["table"]["style_id"] is None)
    check("spacing.list_indent の既定は null",
          u.core.DEFAULT_CONFIG["spacing"]["list_indent"] is None)
    check("table.explicit_format の既定は false",
          u.core.DEFAULT_CONFIG["table"]["explicit_format"] is False)
    check("既定では marL を書き込まない",
          not _has_marl(r2["pptx"]), "marL が書かれている")
    check("明示指定すれば marL を書き込む", _has_marl(r["pptx"]))

    print("[20] 校正言語をテンプレートに合わせ、校正をオフにする")
    # fixtures/template.pptx は python-pptx 既定テンプレート由来で全部 en-US。
    # 設定なしの変換がそれを写すことが「テンプレートを見ている」証拠になる。
    check("設定なしでもスライドに lang が書かれる",
          langs(r2["pptx"], "ppt/slides/") == {"en-US"},
          str(sorted(langs(r2["pptx"], "ppt/slides/"))))
    check("ドキュメント全体が同じ言語になる",
          langs(r2["pptx"]) == {"en-US"}, str(sorted(langs(r2["pptx"]))))
    check("ノートとノートマスターにも入る",
          langs(r2["pptx"], "ppt/notesSlides/", "ppt/notesMasters/") == {"en-US"},
          str(sorted(langs(r2["pptx"], "ppt/notesSlides/", "ppt/notesMasters/"))))
    check("lang を持たない rPr が 1 つも無い",
          unlabeled_rprs(r2["pptx"]) == 0, str(unlabeled_rprs(r2["pptx"])))
    check("既定で noProof がオン", noproofs(r2["pptx"]) == {"1"},
          str(sorted(noproofs(r2["pptx"]))))
    check("language_used に判定結果が返る", r2["language_used"] == "en-US",
          str(r2["language_used"]))

    print("[20b] config で指定した言語が勝つ")
    r20 = u.convert(md, tpl, "t.pptx", "options:\n  language: ja-JP\n")
    check("全パートが ja-JP になる", langs(r20["pptx"]) == {"ja-JP"},
          str(sorted(langs(r20["pptx"]))))
    check("警告なし・枚数も変わらない",
          not r20["warnings"] and r20["slide_count"] == r2["slide_count"],
          str(r20["warnings"]))
    check("fixtures/config.yaml でも ja-JP になる",
          langs(r["pptx"]) == {"ja-JP"}, str(sorted(langs(r["pptx"]))))

    print("[20c] no_proof の 3 値")
    r20c = u.convert(md, tpl, "t.pptx", "options:\n  no_proof: false\n")
    check("false なら noProof=0", noproofs(r20c["pptx"]) == {"0"},
          str(sorted(noproofs(r20c["pptx"]))))
    r20c2 = u.convert(md, tpl, "t.pptx", "options:\n  no_proof: null\n")
    check("null なら noProof を書かない", noproofs(r20c2["pptx"]) == set(),
          str(sorted(noproofs(r20c2["pptx"]))))
    check("null でも lang は書く", langs(r20c2["pptx"], "ppt/slides/") == {"en-US"})

    print("[20d] language: null なら書き込まない")
    r20d = u.convert(md, tpl, "t.pptx", "options:\n  language: null\n")
    check("スライドに lang が無い", langs(r20d["pptx"], "ppt/slides/") == set(),
          str(sorted(langs(r20d["pptx"], "ppt/slides/"))))
    check("レイアウトはテンプレートのまま",
          langs(r20d["pptx"], "ppt/slideLayouts/") == {"en-US"})
    check("language_used は None", r20d["language_used"] is None)
    r20d2 = u.convert(md, tpl, "t.pptx",
                      "options:\n  language: null\n  no_proof: null\n")
    check("両方 null なら lang も noProof も書かない",
          langs(r20d2["pptx"], "ppt/slides/") == set()
          and noproofs(r20d2["pptx"]) == set())

    print("[20e] テンプレートに lang が無ければ日本語にする")
    bare = strip_langs(tpl)
    check("剥がしたテンプレートに lang が無い", langs(bare) == set(),
          str(sorted(langs(bare))))
    r20e = u.convert(md, bare, "t.pptx", None)
    check("フォールバックで ja-JP になる", langs(r20e["pptx"]) == {"ja-JP"},
          str(sorted(langs(r20e["pptx"]))))
    from pptx import Presentation  # noqa: E402
    import io as _io  # noqa: E402
    check("detect_language がテンプレートを読む",
          u.core.detect_language(Presentation(_io.BytesIO(tpl))) == "en-US")
    check("lang が無ければ既定値",
          u.core.detect_language(Presentation(_io.BytesIO(bare)))
          == u.core.DEFAULT_LANGUAGE)

    print("[20f] 既定値と不正な指定")
    check("options.language の既定は auto",
          u.core.DEFAULT_CONFIG["options"]["language"] == "auto")
    check("options.no_proof の既定は true",
          u.core.DEFAULT_CONFIG["options"]["no_proof"] is True)
    check("自動判定が detected に載る",
          u.detect_config(u.analyze_template(tpl, "t.pptx"))["options"]["language"]
          == "en-US")
    r20f = u.convert(md, tpl, "t.pptx", "options:\n  language: 日本語\n")
    check("不正な言語タグは警告を出してテンプレートの言語に戻す",
          langs(r20f["pptx"]) == {"en-US"} and len(r20f["warnings"]) == 1,
          str(r20f["warnings"]))

    print("[21] 文字サイズもテンプレートに任せる")
    check("既定値は 15 項目すべて null",
          set(u.core.DEFAULT_CONFIG["sizes"].values()) == {None},
          str(u.core.DEFAULT_CONFIG["sizes"]))
    # 小見出しも縮小も無く 1 枚に収まる Markdown なら、sz は 1 つも書かれない。
    r21 = u.convert("## A\n\n- x\n", tpl, "t.pptx", None)
    check("最小の原稿では sz を 1 つも書かない", run_sizes(r21["pptx"]) == set(),
          str(sorted(run_sizes(r21["pptx"]))))
    check("継承する run がちゃんとある", unsized_runs(r21["pptx"]) > 0)
    check("既定でも大半の run は sz を持たない", unsized_runs(r2["pptx"]) > 0,
          str(unsized_runs(r2["pptx"])))
    check("明示指定すれば書かれる", 32.0 in run_sizes(r["pptx"]),
          str(sorted(run_sizes(r["pptx"]))))
    check("警告なし", not r2["warnings"], str(r2["warnings"]))

    print("[21b] テンプレートの実効サイズを読む")
    from pptx import Presentation as _Prs  # noqa: E402
    import io as _io2  # noqa: E402
    prs = _Prs(_io2.BytesIO(tpl))
    master = prs.slide_masters[0]
    check("titleStyle lvl1 = 44",
          u.core.master_style_size(master, "titleStyle", 1) == 44.0)
    check("bodyStyle lvl1-5 = 32/28/24/20/20",
          [u.core.master_style_size(master, "bodyStyle", n) for n in range(1, 6)]
          == [32.0, 28.0, 24.0, 20.0, 20.0])
    check("otherStyle lvl1 = 18",
          u.core.master_style_size(master, "otherStyle", 1) == 18.0)
    check("defaultTextStyle = 18", u.core.default_text_size(prs) == 18.0)
    章扉 = next(l for l in master.slide_layouts if l.name == "章扉")
    check("章扉のタイトルは ph 側で 40 に上書きされている",
          u.core.placeholder_size(章扉, 0) == 40.0)
    本文 = next(l for l in master.slide_layouts if l.name == "本文")
    check("本文レイアウトは ph 側の上書きを持たない",
          u.core.placeholder_size(本文, 1) is None)

    print("[21c] 対応物が無い項目はテンプレートの本文サイズから算出する")
    # 手動レイアウトの図形はテンプレートの既定テキスト（18pt）を継承するので、
    # そこから算出した値は現行の既定値とちょうど一致する。
    sizes2 = run_sizes(r2["pptx"])
    check("ページ番号は 10pt", 10.0 in sizes2, str(sorted(sizes2)))
    check("コードは 12pt", 12.0 in sizes2, str(sorted(sizes2)))
    check("小見出しは本文より大きい", 20.0 in sizes2, str(sorted(sizes2)))

    print("[21d] 個別に null / 明示ができる")
    for key in sorted(u.core.DEFAULT_CONFIG["sizes"]):
        try:
            one = u.convert(md, tpl, "t.pptx", f"sizes:\n  {key}: null\n")
            check(f"{key}: null で変換できる", not one["warnings"], str(one["warnings"]))
        except Exception as e:                                  # noqa: BLE001
            check(f"{key}: null で変換できる", False, f"{type(e).__name__}: {e}")
    for body in ("null", "[18, null, 14]", "18", "[]"):
        rb = u.convert(md, tpl, "t.pptx", f"sizes:\n  body: {body}\n")
        check(f"body: {body} が通る", not rb["warnings"], str(rb["warnings"]))
    rq = u.convert(md, tpl, "t.pptx", 'sizes:\n  quote: ""\n')
    check("空文字も未指定扱い", not rq["warnings"], str(rq["warnings"]))
    rt = u.convert(md, tpl, "t.pptx", "sizes:\n  table: 9\n")
    check("明示した値が書かれる", 9.0 in run_sizes(rt["pptx"]),
          str(sorted(run_sizes(rt["pptx"]))))

    print("[21e] テンプレートに sz が無ければ組み込みの既定に落ちる")
    bare_sz = strip_sizes(tpl)
    r21e = u.convert(md, bare_sz, "t.pptx", None)
    check("sz を剥がしても変換できる", not r21e["warnings"], str(r21e["warnings"]))
    check("フォールバック（18pt 基準）で描かれる",
          {10.0, 12.0} <= run_sizes(r21e["pptx"]),
          str(sorted(run_sizes(r21e["pptx"]))))

    print("[21f] 不正な値は分かりやすく落とす")
    for label, bad in (
        ("文字列", "sizes:\n  title: 大きめ\n"),
        ("範囲外", "sizes:\n  title: 0\n"),
        ("真偽値", "sizes:\n  title: yes\n"),
        ("タイポ", "sizes:\n  titel: 32\n"),
        ("リストの中の文字列", "sizes:\n  body: [18, 大]\n"),
    ):
        try:
            u.convert(md, tpl, "t.pptx", bad)
            check(f"{label} はエラー", False, "(no exception)")
        except u.Md2pptError as e:
            check(f"{label} はエラー", True)
            print(f"       -> {e}")

    print("[21g] fonts / colors / sizes の null を併用できる")
    r21g = u.convert(md, tpl, "t.pptx", no_fonts + no_colors)
    check("3 つとも既定 null で変換できる",
          typefaces(r21g["pptx"]) == set() and not r21g["warnings"],
          str(r21g["warnings"]))

    print()
    if failures:
        print(f"FAILED: {len(failures)} check(s): {', '.join(failures)}")
        return 1
    print("all checks passed")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
