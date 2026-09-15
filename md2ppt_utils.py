"""Dify 向けの薄いラッパ。

md2pptx_core は sample/md2pptx.py の無改造コピーで、パス入出力と `SystemExit` を
前提にした CLI 向けの作りになっている。ここではその差分だけを吸収して、
バイト列 in / バイト列 out のインタフェースを提供する。

このモジュールは dify_plugin を import しない（アップロードファイルは
`.blob` / `.filename` をダックタイピングで読む）。SDK 無しで単体テストできる。
"""

from __future__ import annotations

import copy
import io
import re
import tempfile
import zipfile
from pathlib import Path
from typing import Any

import yaml
from pptx import Presentation

import md2pptx_core as core

EMU_IN = 914400.0

TEMPLATE_SUFFIXES = (".pptx", ".potx")

PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

# .potx は [Content_Types].xml で template.main+xml を宣言する。python-pptx の
# Presentation() は presentation.main+xml しか受け付けないため、拡張子を .pptx に
# 変えるだけでは開けない（md2pptx_core._open_template のやり方では不十分）。
_POTX_CT = "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"
_PPTX_CT = "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"


class Md2pptError(Exception):
    """利用者にそのまま見せられるエラー。"""


# ════════════════════════════════════════════════════════════════════════
# 入力パラメータの取り出し（SDK 非依存）
# ════════════════════════════════════════════════════════════════════════

# Dify が未入力の任意パラメータに入れて送ってくることがある番兵値。
# SDK は値を一切加工せず素通しするため、サーバ側の str(None) が
# そのまま文字列 "None" として届く。YAML の null キーワード（null / ~）は
# safe_load が None にしてくれるが、"None" はただの文字列として通ってしまう。
_UNSET_TEXT = frozenset({"none", "null", "nil", "undefined", "~", "-"})

_TRUE_TEXT = frozenset({"true", "yes", "on", "1"})
_FALSE_TEXT = frozenset({"false", "no", "off", "0", "none", "null", ""})


def text_param(value: Any) -> str:
    """任意の文字列パラメータを正規化する。未入力は "" にする。

    これらの番兵値を意味のある入力として扱う利用者は想定できないため、
    未入力とみなして握りつぶすほうが実害が小さい。
    """
    if value is None:
        return ""
    if not isinstance(value, str):
        value = str(value)
    text = value.strip()
    return "" if text.lower() in _UNSET_TEXT else text


def bool_param(value: Any, default: bool = False) -> bool:
    """真偽パラメータを正規化する。

    bool("false") が True になる罠を避ける。Dify が真偽値を文字列で
    送ってきても意図どおりに解釈する。
    """
    if value is None:
        return default
    if isinstance(value, bool):
        return value
    if isinstance(value, str):
        low = value.strip().lower()
        if low in _TRUE_TEXT:
            return True
        if low in _FALSE_TEXT:
            return False
        return default
    return bool(value)


def file_bytes(param: Any) -> bytes | None:
    """dify_plugin の File / dict / bytes からファイル内容を取り出す。"""
    if param is None or param == "":
        return None
    if isinstance(param, (bytes, bytearray)):
        return bytes(param)
    if isinstance(param, (list, tuple)):
        raise Md2pptError("単一ファイルのパラメータに複数のファイルが渡されました。")

    blob = getattr(param, "blob", None)
    if blob is None and isinstance(param, dict):
        blob = param.get("blob")
    if blob is None:
        raise Md2pptError(
            "アップロードされたファイルの内容を取得できませんでした。"
            "リモートデバッグ中の場合は Dify 側の FILES_URL が"
            "このプラグインから到達可能な URL になっているか確認してください。"
        )
    return bytes(blob)


def file_name_of(param: Any, default: str = "") -> str:
    """File / dict からファイル名を取り出す。"""
    name = getattr(param, "filename", None)
    if not name and isinstance(param, dict):
        name = param.get("filename") or param.get("file_name")
    return str(name or default)


def decode_text(data: bytes) -> str:
    """UTF-8 を基本に、日本語環境でありがちな符号化にフォールバックする。"""
    for enc in ("utf-8-sig", "utf-8", "cp932", "utf-16"):
        try:
            return data.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    raise Md2pptError(
        "ファイルの文字コードを判別できませんでした。UTF-8 で保存し直してください。"
    )


def read_markdown(text_value: Any, file_value: Any) -> tuple[str, bool]:
    """Markdown 本文と、ファイル入力を採用したかどうかを返す。"""
    data = file_bytes(file_value)
    if data is not None:
        text = decode_text(data)
        used_file = True
    else:
        text = text_param(text_value)
        used_file = False
    if not text.strip():
        raise Md2pptError(
            "Markdown が空です。markdown_text か markdown_file のどちらかを指定してください。"
        )
    return text.replace("\r\n", "\n").replace("\r", "\n"), used_file


def ensure_pptx(name: Any) -> str:
    """出力ファイル名を安全な .pptx 名に整える。"""
    cleaned = re.sub(r'[\\/:*?"<>|\r\n\t]', "_", text_param(name)).strip(". ")
    if not cleaned:
        cleaned = "presentation"
    return cleaned if cleaned.lower().endswith(".pptx") else f"{cleaned}.pptx"


# ════════════════════════════════════════════════════════════════════════
# 設定
# ════════════════════════════════════════════════════════════════════════

# DEFAULT_CONFIG のトップレベルキーのうち、マッピングを期待するもの。
_CONFIG_SECTIONS = (
    "layouts", "placeholders", "body_area", "fonts", "sizes", "colors",
    "table", "quote", "image", "spacing", "options",
)
# テンプレートは毎回アップロードするので、config 側の指定は黙って無視する。
_CONFIG_IGNORED = ("template",)


def _fullwidth_hint(text: str) -> str:
    """日本語環境でありがちな全角文字の混入を具体的に指摘する。"""
    if "：" in text:
        return "（全角コロン「：」が含まれています。半角の「:」にしてください）"
    if "　" in text:
        return "（全角スペースが含まれています。字下げは半角スペースにしてください）"
    return ""


def parse_config_yaml(config_yaml: Any) -> dict:
    """config_yaml パラメータを辞書にする。未指定なら空の辞書。"""
    if isinstance(config_yaml, dict):
        return config_yaml          # 既に辞書ならそのまま使う
    if isinstance(config_yaml, (bytes, bytearray)):
        config_yaml = decode_text(bytes(config_yaml))

    text = text_param(config_yaml)
    if not text:
        return {}
    # 設定として成立する最低条件はコロンを含むこと。コロンも改行も無い単一
    # トークンは設定になりえないので未入力として扱う。_UNSET_TEXT の列挙だけでは
    # Dify が未入力欄に別の文字列を入れて送ってきたときにまた止まってしまう。
    # 全角コロンはここで握りつぶさず、下の _fullwidth_hint による指摘へ回す。
    if "\n" not in text and ":" not in text and "：" not in text:
        return {}

    try:
        data = yaml.safe_load(text)
    except yaml.YAMLError as e:
        raise Md2pptError(
            f"config_yaml を解釈できません{_fullwidth_hint(text)}: {e}"
        ) from None
    if data is None:
        return {}
    if not isinstance(data, dict):
        raise Md2pptError(
            "config_yaml は 'layouts:' のようなマッピング（key: value）で"
            f"記述してください{_fullwidth_hint(text)}。"
            f"受け取った値（{type(data).__name__}）: {text[:80]!r}"
        )

    _validate_config(data, text)
    return data


def _validate_config(data: dict, text: str) -> None:
    """後段で分かりにくく失敗する前に、設定の形を確かめる。"""
    unknown = [
        k for k in data
        if k not in _CONFIG_SECTIONS and k not in _CONFIG_IGNORED
    ]
    if unknown:
        # 全角スペースで字下げすると、その行がトップレベルの
        # 「　　content」のようなキーとして拾われる。
        raise Md2pptError(
            f"config_yaml に不明なキーがあります: {', '.join(map(str, unknown))}"
            f"{_fullwidth_hint(text)}。"
            f"指定できるのは {', '.join(_CONFIG_SECTIONS)} です。"
        )
    for key in _CONFIG_SECTIONS:
        if key in data and not isinstance(data[key], dict):
            # 全角スペースで字下げすると {'layouts': None, '　　content': ...} になり、
            # そのまま進むと Renderer.layout() で AttributeError になる。
            raise Md2pptError(
                f"config_yaml の '{key}' はマッピング（key: value）で"
                f"記述してください{_fullwidth_hint(text)}。"
                f"受け取った値（{type(data[key]).__name__}）: {str(data[key])[:60]!r}"
            )
    _validate_sizes(data, text)


def _validate_sizes(data: dict, text: str) -> None:
    """sizes の値が数値か null かを確かめる。

    core 側の opt_size() も不正値を「指定なし」に倒して自衛するが、Dify 経由の
    場合はここで具体的に指摘したほうが直しやすい。
    """
    sizes = data.get("sizes") or {}
    known = set(core.DEFAULT_CONFIG["sizes"])
    unknown = [k for k in sizes if k not in known]
    if unknown:
        raise Md2pptError(
            f"config_yaml の sizes に不明なキーがあります: "
            f"{', '.join(map(str, unknown))}{_fullwidth_hint(text)}。"
            f"指定できるのは {', '.join(sorted(known))} です。"
        )
    for key, value in sizes.items():
        # body だけはレベル別のリストを書ける。
        items = value if key == "body" and isinstance(value, list) else [value]
        for item in items:
            if item is None or (isinstance(item, str) and not item.strip()):
                continue                     # null / 空欄は「指定なし」
            # bool は int の一種なので、yes / no を書かれると素通りしてしまう。
            if isinstance(item, bool) or not _is_number(item):
                raise Md2pptError(
                    f"config_yaml の sizes.{key} は数値か null で"
                    f"指定してください{_fullwidth_hint(text)}。"
                    f"受け取った値（{type(item).__name__}）: {str(item)[:40]!r}"
                )
            if not core.MIN_FONT_SIZE <= float(item) <= core.MAX_FONT_SIZE:
                raise Md2pptError(
                    f"config_yaml の sizes.{key} は "
                    f"{core.MIN_FONT_SIZE:g}〜{core.MAX_FONT_SIZE:g} の範囲で"
                    f"指定してください（受け取った値: {item}）。"
                )


def _is_number(value) -> bool:
    """数値、または数値として読める文字列か。"""
    if isinstance(value, (int, float)):
        return True
    if isinstance(value, str):
        try:
            float(value.strip())
            return True
        except ValueError:
            return False
    return False


def build_config(config_yaml: str | None, detected: dict | None = None) -> dict:
    """設定を組み立てる。

    優先順位は DEFAULT_CONFIG < detected（テンプレートからの自動判定） < config_yaml。
    利用者が明示した項目は常に勝つので、自動判定が邪魔をすることはない。
    """
    cfg = copy.deepcopy(core.DEFAULT_CONFIG)
    if detected:
        cfg = core.deep_merge(cfg, detected)
    user = parse_config_yaml(config_yaml)
    if user:
        cfg = core.deep_merge(cfg, user)
    return cfg


# ════════════════════════════════════════════════════════════════════════
# テンプレート
# ════════════════════════════════════════════════════════════════════════

def check_template_name(filename: str | None) -> str:
    """拡張子を検証して小文字で返す。"""
    suffix = Path(filename or "template.pptx").suffix.lower()
    if suffix not in TEMPLATE_SUFFIXES:
        raise Md2pptError(
            f"テンプレートは .pptx または .potx を指定してください"
            f"（受け取った拡張子: {suffix or 'なし'}）"
        )
    return suffix


def prepare_template(template_bytes: bytes, workdir: Path) -> Path:
    """テンプレートを workdir に .pptx として書き出し、絶対パスを返す。

    - 常に `.pptx` 名で置くので core._open_template の .potx 分岐
      （後始末されない tempfile.mkdtemp を使う）を通らない。
    - .potx の content type はここで .pptx 相当に書き換える。
    - 絶対パスなので core._find_template が即 return し、Path.cwd() を見ない。
    """
    if not template_bytes:
        raise Md2pptError("テンプレートファイルが空です。")
    if template_bytes[:2] != b"PK":
        raise Md2pptError(
            "テンプレートが Office Open XML（.pptx / .potx）ではありません。"
            "古い .ppt 形式の場合は PowerPoint で .pptx に保存し直してください。"
        )
    dst = workdir / "template.pptx"
    dst.write_bytes(template_bytes)
    _normalize_content_type(dst)
    return dst.resolve()


def _normalize_content_type(path: Path) -> None:
    """.potx の content type を .pptx 相当に書き換える。.pptx なら何もしない。"""
    try:
        with zipfile.ZipFile(path) as zf:
            try:
                ct = zf.read("[Content_Types].xml").decode("utf-8")
            except KeyError:
                raise Md2pptError(
                    "テンプレートの構造が不正です（[Content_Types].xml がありません）。"
                ) from None
            if _POTX_CT not in ct:
                return
            entries = [(info, zf.read(info.filename)) for info in zf.infolist()]
    except zipfile.BadZipFile:
        raise Md2pptError("テンプレートを展開できません。ファイルが破損しています。") from None

    tmp = path.with_name(path.name + ".fixed")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as out:
        for info, data in entries:
            if info.filename == "[Content_Types].xml":
                data = ct.replace(_POTX_CT, _PPTX_CT).encode("utf-8")
            out.writestr(info, data)
    tmp.replace(path)


def _open_presentation(path: Path) -> Presentation:
    try:
        return Presentation(path)
    except Md2pptError:
        raise
    except Exception as e:
        raise Md2pptError(
            f"テンプレートを開けませんでした。破損しているか "
            f"PowerPoint 形式ではない可能性があります: {e}"
        ) from None


# ════════════════════════════════════════════════════════════════════════
# 変換
# ════════════════════════════════════════════════════════════════════════

def convert(
    md_text: str,
    template_bytes: bytes,
    template_filename: str | None = None,
    config_yaml: str | None = None,
) -> dict[str, Any]:
    """Markdown を PowerPoint に変換する。

    returns: {"pptx", "slide_count", "spec_count", "outline", "warnings",
              "meta", "detected"}
    """
    check_template_name(template_filename)
    # config_yaml が壊れているならテンプレートを開く前に落とす
    parse_config_yaml(config_yaml)

    with tempfile.TemporaryDirectory(prefix="md2ppt-") as td:
        work = Path(td)
        path = prepare_template(template_bytes, work)

        # レイアウト名は既定値（表紙 / 章扉 / 本文 …）とは限らないので、
        # テンプレートの中身から実際のレイアウトを判定して既定値の上に敷く。
        # config_yaml で明示された項目は引き続き最優先される。
        info = describe_presentation(
            _open_presentation(path), Path(template_filename or "template.pptx").name
        )
        detected = detect_config(info)

        cfg = build_config(config_yaml, detected)
        cfg["template"] = str(path)

        meta, blocks, notes = core.parse_markdown(md_text)
        specs = core.build_slides(meta, blocks, cfg)
        if not specs:
            raise Md2pptError(
                "スライドを 1 枚も作れませんでした。"
                "見出し（# / ##）か front matter の title を含めてください。"
            )
        specs[0].notes.extend(notes[:1])   # core.main() と同じ挙動

        try:
            # base_dir は一時ディレクトリ。相対パスの画像はここに解決される。
            renderer = core.Renderer(cfg, work, work)
            renderer.render(specs)
        except Md2pptError:
            raise
        except SystemExit as e:
            # core._open_template はテンプレート不在時に SystemExit を投げる。
            # BaseException なので except Exception では捕まらない。
            raise Md2pptError(str(e) or "テンプレートの読み込みに失敗しました") from None
        except Exception as e:
            raise Md2pptError(f"PowerPoint の生成に失敗しました: {e}") from None

        buf = io.BytesIO()
        renderer.prs.save(buf)   # python-pptx はファイルライクを受け付ける

        return {
            "pptx": buf.getvalue(),
            "slide_count": len(renderer.prs.slides),
            "spec_count": len(specs),
            "outline": core.outline(specs),
            "warnings": list(dict.fromkeys(renderer.warnings)),
            "meta": {
                k: str(v) for k, v in (meta or {}).items()
                if k in ("title", "subtitle", "author", "date") and v
            },
            "detected": detected,
            "layouts_used": dict(cfg["layouts"]),
            "language_used": renderer.applied_language,
        }


# ════════════════════════════════════════════════════════════════════════
# テンプレート解析（sample/inspect_template.py 相当）
# ════════════════════════════════════════════════════════════════════════

def _inch(v) -> float | None:
    return None if v is None else round(v / EMU_IN, 2)


def describe_presentation(prs: Presentation, name: str) -> dict[str, Any]:
    """開いた Presentation からレイアウトとプレースホルダの一覧を作る。"""
    layouts = []
    for i, layout in enumerate(prs.slide_layouts):
        placeholders = []
        for ph in layout.placeholders:
            pf = ph.placeholder_format
            placeholders.append({
                "idx": pf.idx,
                "type": str(pf.type).split(" ")[0],
                "name": ph.name,
                "left_in": _inch(ph.left),
                "top_in": _inch(ph.top),
                "width_in": _inch(ph.width),
                "height_in": _inch(ph.height),
            })
        layouts.append({
            "index": i,
            "name": layout.name,
            "placeholders": placeholders,
            "other_shapes": [s.name for s in layout.shapes if not s.is_placeholder],
        })

    w_in = _inch(prs.slide_width) or 0.0
    h_in = _inch(prs.slide_height) or 0.0
    return {
        "template_name": name,
        "slide_width_in": w_in,
        "slide_height_in": h_in,
        "aspect_ratio": _aspect(w_in, h_in),
        "master_count": len(prs.slide_masters),
        "layout_count": len(prs.slide_layouts),
        "layout_names": [l["name"] for l in layouts],
        "layouts": layouts,
        # テンプレートに書かれている校正言語（options.language: auto の既定値）
        "language": core.detect_language(prs),
        "language_counts": core.count_languages(prs),
    }


def analyze_template(
    template_bytes: bytes, template_filename: str | None = None
) -> dict[str, Any]:
    """レイアウトとプレースホルダの一覧を辞書で返す。"""
    check_template_name(template_filename)
    name = Path(template_filename or "template.pptx").name

    with tempfile.TemporaryDirectory(prefix="md2ppt-") as td:
        path = prepare_template(template_bytes, Path(td))
        return describe_presentation(_open_presentation(path), name)


def _aspect(w: float, h: float) -> str:
    if not w or not h:
        return "-"
    r = w / h
    if abs(r - 16 / 9) < 0.02:
        return "16:9"
    if abs(r - 4 / 3) < 0.02:
        return "4:3"
    return f"{r:.2f}:1"


def _language_detail(info: dict[str, Any]) -> str:
    """校正言語の内訳。テンプレートに lang が無ければ既定値だと明示する。"""
    counts = info.get("language_counts") or {}
    if not counts:
        return f"（テンプレートに lang が無いため既定の {core.DEFAULT_LANGUAGE}）"
    inner = ", ".join(f"{k} {v}" for k, v in
                      sorted(counts.items(), key=lambda kv: (-kv[1], kv[0])))
    return f"（{inner}）"


def format_template_report(info: dict[str, Any]) -> str:
    """inspect_template.describe() と同じ体裁の人間可読テキスト。"""

    def num(v) -> str:
        return "-" if v is None else f"{v:.2f}"

    out = [
        f"# {info['template_name']}",
        f"スライドサイズ: {num(info['slide_width_in'])}in x "
        f"{num(info['slide_height_in'])}in ({info['aspect_ratio']})",
        f"スライドマスタ数: {info['master_count']} / レイアウト数: {info['layout_count']}",
        f"校正言語: {info['language']}{_language_detail(info)}",
        "",
    ]
    for layout in info["layouts"]:
        out.append(f"[{layout['index']}] {layout['name']!r}")
        out.append(
            "      idx  type                 name                      "
            "left   top    width  height  (inch)"
        )
        for ph in layout["placeholders"]:
            out.append(
                f"      {ph['idx']:<4} {ph['type']:<20} {ph['name'][:24]:<25} "
                f"{num(ph['left_in']):>6} {num(ph['top_in']):>6} "
                f"{num(ph['width_in']):>6} {num(ph['height_in']):>6}"
            )
        if layout["other_shapes"]:
            out.append(
                f"      (プレースホルダ以外の図形: {', '.join(layout['other_shapes'])})"
            )
        out.append("")
    return "\n".join(out)


_GUESS = {
    "title": ("表紙", "title slide", "タイトル スライド", "cover", "扉"),
    "section": ("章扉", "section", "セクション", "章見出し"),
    "content": ("本文", "title and content", "タイトルとコンテンツ", "コンテンツ"),
    "table": ("タイトルのみ", "title only"),
    "blank": ("白紙", "blank"),
}

# プレースホルダの分類。日付・フッター・ページ番号は「中身」とは数えない。
_TITLE_PH = ("TITLE", "CENTER_TITLE", "VERTICAL_TITLE")
_CHROME_PH = ("DATE", "FOOTER", "SLIDE_NUMBER")
_BODY_PH = ("BODY", "OBJECT", "VERTICAL_BODY", "VERTICAL_OBJECT")


def _title_placeholder(layout: dict) -> dict | None:
    return next((p for p in layout["placeholders"] if p["type"] in _TITLE_PH), None)


def _content_placeholders(layout: dict) -> list[dict]:
    return [
        p for p in layout["placeholders"]
        if p["type"] not in _TITLE_PH and p["type"] not in _CHROME_PH
    ]


def detect_config(info: dict[str, Any]) -> dict[str, Any]:
    """テンプレートの中身から layouts / placeholders を判定する。

    まずレイアウト名のキーワードで探し、見つからなければプレースホルダの
    構成から推測する。名前が英語でも独自の命名でも、標準的なレイアウト構成の
    テンプレートであれば設定なしで動く。
    """
    layouts = info["layouts"]
    if not layouts:
        return {}

    def by_name(key: str) -> dict | None:
        keys = _GUESS[key]
        for layout in layouts:
            low = layout["name"].lower()
            if any(k.lower() in low for k in keys):
                return layout
        return None

    def first(pred) -> dict | None:
        return next((l for l in layouts if pred(l)), None)

    # 表紙: SUBTITLE を持つ、なければ CENTER_TITLE を持つレイアウト
    cover = (
        by_name("title")
        or first(lambda l: any(p["type"] == "SUBTITLE" for p in l["placeholders"]))
        or first(lambda l: any(p["type"] == "CENTER_TITLE" for p in l["placeholders"]))
        or layouts[0]
    )

    # 本文: タイトル＋本文プレースホルダがちょうど 1 つ（2 カラムや比較を避ける）
    content = (
        by_name("content")
        or first(lambda l: _title_placeholder(l)
                 and len(_content_placeholders(l)) == 1
                 and _content_placeholders(l)[0]["type"] in _BODY_PH)
        or first(lambda l: _title_placeholder(l) and _content_placeholders(l))
        or layouts[0]
    )

    # タイトルのみ: タイトルはあるが本文プレースホルダが無い（手動レイアウト用）
    only = (
        by_name("table")
        or first(lambda l: _title_placeholder(l) and not _content_placeholders(l))
    )

    # 白紙: プレースホルダが（装飾を除いて）無い
    blank = (
        by_name("blank")
        or first(lambda l: not _title_placeholder(l) and not _content_placeholders(l))
        or only
        or content
    )

    # 章扉: 大きなタイトルだけ置ければよいので、タイトルのみ → 本文の順に代用
    section = by_name("section") or only or content

    title_ph = _title_placeholder(content)
    body_ph = next(
        (p for p in _content_placeholders(content) if p["type"] in _BODY_PH), None
    )
    sub_ph = next(
        (p for p in _content_placeholders(cover)
         if p["type"] in ("SUBTITLE",) + _BODY_PH), None
    )

    return {
        "layouts": {
            "title": cover["name"],
            "section": section["name"],
            "content": content["name"],
            "table": (only or blank or content)["name"],
            "blank": blank["name"],
        },
        "placeholders": {
            "title": title_ph["idx"] if title_ph else 0,
            "subtitle": sub_ph["idx"] if sub_ph else 1,
            "body": body_ph["idx"] if body_ph else 1,
        },
        # 既定の "auto" を、テンプレートから読み取った実際の言語に置き換える。
        "options": {"language": info["language"]},
    }


def config_to_yaml(suggested: dict[str, Any], template_name: str) -> str:
    """detect_config() の結果を config_yaml パラメータに貼れる YAML にする。"""
    body = yaml.safe_dump(
        suggested, allow_unicode=True, sort_keys=False, default_flow_style=False
    )
    return (
        f"# md2ppt 設定（inspect_template ツールが {template_name} から生成）\n"
        f"# md_to_pptx ツールの config_yaml パラメータにこのまま貼り付けてください。\n"
        f"{body}"
    )


# ════════════════════════════════════════════════════════════════════════
# ツールの出力
# ════════════════════════════════════════════════════════════════════════

def payload_messages(tool: Any, payload: dict[str, Any]):
    """JSON ペイロードと、その全キーのワークフロー変数を yield する。

    ツール定義の output_schema は Dify の変数ピッカーに名前と型を宣言するだけで、
    値は入らない。後続のノードから参照できるようにするには、変数メッセージを
    別に流す必要がある。同じ dict から JSON と変数の両方を作ることで、
    「JSON には値があるのに変数は null」という食い違いが起きないようにする。

    値が None のキーも null のまま出す。ピッカーに出る名前の集合と実際に出る
    変数の集合を常に一致させておくためで、テスト（[9c]）がこれを見張っている。
    """
    yield tool.create_json_message(payload)
    for name, value in payload.items():
        yield tool.create_variable_message(name, value)
